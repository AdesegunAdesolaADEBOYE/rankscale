import csv
import io
import json
import logging
from datetime import date, datetime

from celery import shared_task
from django.db import transaction
from openpyxl import load_workbook
from pypdf import PdfReader

from core.google_sheets import (
    GoogleSheetsError,
    credentials_for_connection,
    download_public_sheet_csv,
    iter_private_sheet_rows,
)
from core.models import Dataset, DatasetRow, GoogleSheetsConnection

logger = logging.getLogger(__name__)

MAX_COLUMNS = 50
MAX_ROWS = 100_000
MAX_PDF_PAGES = 1_000
ROW_BATCH_SIZE = 500
MAX_CELL_SIZE = 1_000_000


class DatasetImportError(Exception):
    pass


def _clean_columns(values):
    columns = [str(value).strip() for value in values]
    if not columns or any(not column for column in columns):
        raise DatasetImportError("Every column needs a header.")
    if len(columns) > MAX_COLUMNS:
        raise DatasetImportError(f"A dataset can have at most {MAX_COLUMNS} columns.")
    if len({column.casefold() for column in columns}) != len(columns):
        raise DatasetImportError("Column headers must be unique.")
    return columns


def _value_size(value):
    if isinstance(value, str):
        return len(value)
    return len(json.dumps(value, ensure_ascii=False, separators=(",", ":")))


def _save_rows(dataset, columns, records):
    columns = _clean_columns(columns)
    batch = []
    row_count = 0

    for row_number, values in records:
        if isinstance(values, dict):
            if set(values) != set(columns):
                raise DatasetImportError(
                    f"Row {row_number} has different columns from the header."
                )
            data = {column: values[column] for column in columns}
        else:
            values = list(values)
            if not values or not any(value not in (None, "") for value in values):
                continue
            if len(values) != len(columns):
                raise DatasetImportError(
                    f"Row {row_number} has {len(values)} values; expected {len(columns)}."
                )
            data = dict(zip(columns, values))

        if all(value in (None, "") for value in data.values()):
            continue
        if any(_value_size(value) > MAX_CELL_SIZE for value in data.values()):
            raise DatasetImportError(f"Row {row_number} contains a value that is too large.")
        if row_count >= MAX_ROWS:
            raise DatasetImportError(f"A dataset can contain at most {MAX_ROWS:,} data rows.")

        batch.append(DatasetRow(dataset=dataset, row_number=row_number, data=data))
        row_count += 1
        if len(batch) >= ROW_BATCH_SIZE:
            DatasetRow.objects.bulk_create(batch, batch_size=ROW_BATCH_SIZE)
            batch.clear()

    if batch:
        DatasetRow.objects.bulk_create(batch, batch_size=ROW_BATCH_SIZE)
    if row_count == 0:
        raise DatasetImportError("Add at least one data row below the headers.")
    return columns, row_count


def _process_csv(dataset, binary_source):
    with io.TextIOWrapper(binary_source, encoding="utf-8-sig", newline="") as source:
        reader = csv.reader(source, strict=True)
        try:
            columns = _clean_columns(next(reader))
        except StopIteration as exc:
            raise DatasetImportError("This file is empty.") from exc

        def records():
            for row_number, values in enumerate(reader, start=2):
                if not values or not any(value.strip() for value in values):
                    continue
                yield row_number, values

        return _save_rows(dataset, columns, records())


def _excel_value(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def _process_xlsx(dataset, binary_source):
    try:
        workbook = load_workbook(binary_source, read_only=True, data_only=True)
    except Exception as exc:
        raise DatasetImportError("This is not a readable .xlsx workbook.") from exc
    try:
        if not workbook.worksheets:
            raise DatasetImportError("This workbook has no worksheets.")
        worksheet = workbook.worksheets[0]
        rows = worksheet.iter_rows(values_only=True)
        try:
            columns = _clean_columns(_excel_value(value) for value in next(rows))
        except StopIteration as exc:
            raise DatasetImportError("The first worksheet is empty.") from exc

        def records():
            for row_number, raw_values in enumerate(rows, start=2):
                values = [_excel_value(value) for value in raw_values]
                while values and not values[-1]:
                    values.pop()
                if not values or not any(values):
                    continue
                if len(values) > len(columns):
                    raise DatasetImportError(
                        f"Row {row_number} has values beyond the header columns."
                    )
                values.extend([""] * (len(columns) - len(values)))
                yield row_number, values

        return _save_rows(dataset, columns, records())
    finally:
        workbook.close()


def _process_json(dataset, binary_source):
    try:
        with io.TextIOWrapper(binary_source, encoding="utf-8-sig") as source:
            items = json.load(source)
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise DatasetImportError("This file is not valid UTF-8 JSON.") from exc
    if not isinstance(items, list) or not items or not isinstance(items[0], dict):
        raise DatasetImportError("JSON must contain a list of row objects.")

    columns = _clean_columns(items[0].keys())

    def records():
        for row_number, item in enumerate(items, start=1):
            if not isinstance(item, dict) or set(item) != set(columns):
                raise DatasetImportError(
                    f"JSON row {row_number} must have the same keys as the first row."
                )
            yield row_number, item

    return _save_rows(dataset, columns, records())


def _process_pdf(dataset, binary_source):
    try:
        reader = PdfReader(binary_source, strict=True)
        if reader.is_encrypted:
            raise DatasetImportError("Password-protected PDFs are not supported.")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise DatasetImportError(f"A PDF can have at most {MAX_PDF_PAGES:,} pages.")

        def records():
            for page_number, page in enumerate(reader.pages, start=1):
                text = page.extract_text() or ""
                if text.strip():
                    yield page_number, {"page": page_number, "text": text.strip()}

        return _save_rows(dataset, ("page", "text"), records())
    except DatasetImportError:
        raise
    except Exception as exc:
        raise DatasetImportError(
            "This PDF could not be read. Password-protected files and scanned images without text are not supported."
        ) from exc


def _process_google_sheet(dataset):
    try:
        connection = dataset.campaign.owner.google_sheets_connection
    except GoogleSheetsConnection.DoesNotExist:
        connection = None

    if connection:
        credentials = credentials_for_connection(connection)
        rows = iter_private_sheet_rows(dataset.source_url, credentials)
        try:
            _header_row, header_values = next(rows)
        except StopIteration as exc:
            raise DatasetImportError("This Google Sheet is empty.") from exc
        columns = _clean_columns(header_values)

        def records():
            for row_number, values in rows:
                if not values or not any(value not in (None, "") for value in values):
                    continue
                if len(values) > len(columns):
                    raise DatasetImportError(
                        f"Row {row_number} has values beyond the header columns."
                    )
                values = list(values) + [""] * (len(columns) - len(values))
                yield row_number, values

        return _save_rows(dataset, columns, records())

    try:
        content = download_public_sheet_csv(dataset.source_url)
    except GoogleSheetsError as exc:
        if "not publicly viewable" in str(exc):
            raise DatasetImportError(
                "This sheet is private. Connect a Google account with access, then retry."
            ) from exc
        raise DatasetImportError(str(exc)) from exc
    return _process_csv(dataset, io.BytesIO(content))


@shared_task(name="core.process_dataset")
def process_dataset(dataset_id):
    """Parse an uploaded file or Google Sheet into campaign-owned rows."""
    updated = Dataset.objects.filter(
        pk=dataset_id,
        status=Dataset.Status.QUEUED,
    ).update(status=Dataset.Status.PROCESSING, error_message="")
    if not updated:
        return {"dataset_id": dataset_id, "status": "not_queued"}

    dataset = Dataset.objects.select_related("campaign__owner").get(pk=dataset_id)
    dataset.rows.all().delete()

    try:
        if dataset.source_format == Dataset.SourceFormat.GOOGLE_SHEET:
            columns, row_count = _process_google_sheet(dataset)
        else:
            with dataset.file.open("rb") as source:
                processors = {
                    Dataset.SourceFormat.CSV: _process_csv,
                    Dataset.SourceFormat.XLSX: _process_xlsx,
                    Dataset.SourceFormat.JSON: _process_json,
                    Dataset.SourceFormat.PDF: _process_pdf,
                }
                processor = processors.get(dataset.source_format)
                if not processor:
                    raise DatasetImportError("This file format is not supported.")
                columns, row_count = processor(dataset, source)
    except (DatasetImportError, GoogleSheetsError, csv.Error, UnicodeDecodeError) as exc:
        dataset.rows.all().delete()
        dataset.status = Dataset.Status.FAILED
        dataset.row_count = 0
        dataset.columns = []
        dataset.error_message = str(exc)[:500]
        dataset.save(
            update_fields=("status", "row_count", "columns", "error_message", "updated_at")
        )
        return {"dataset_id": dataset_id, "status": "failed"}
    except Exception:
        logger.exception("Unexpected error while importing dataset %s", dataset_id)
        dataset.rows.all().delete()
        dataset.status = Dataset.Status.FAILED
        dataset.row_count = 0
        dataset.columns = []
        dataset.error_message = (
            "We couldn't process this file. Check the selected format and file contents, then try again."
        )
        dataset.save(
            update_fields=("status", "row_count", "columns", "error_message", "updated_at")
        )
        return {"dataset_id": dataset_id, "status": "failed"}

    with transaction.atomic():
        Dataset.objects.filter(
            pk=dataset_id,
            status=Dataset.Status.PROCESSING,
        ).update(
            status=Dataset.Status.READY,
            columns=columns,
            row_count=row_count,
            error_message="",
        )

    return {"dataset_id": dataset_id, "status": "ready", "row_count": row_count}
